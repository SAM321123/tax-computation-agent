"""Format-agnostic document intake for the tax computation agent.

The client document is intentionally not required to follow one workbook
schema. This module profiles the uploaded files, discovers candidate tables or
amount-bearing text, and produces provisional ledger rows for the downstream
agent. In production, ``DocumentUnderstandingAgent`` is the seam where an LLM
or document-intelligence service should replace the local fallback while
keeping the same normalized row contract.
"""
from __future__ import annotations

import json
import math
import re
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Protocol
from xml.etree import ElementTree

import pandas as pd
from openpyxl import load_workbook

from app.intake.trial_balance import categorize
from app.rules_engine.models import LedgerLine


class DocumentUnderstandingAgent(Protocol):
    def inspect(self, path: Path) -> "DocumentProfile": ...


@dataclass
class DocumentProfile:
    filename: str
    extension: str
    document_type: str
    extraction_status: str
    detected_rows: int = 0
    detected_sheets: list[str] | None = None
    text_preview: str = ""
    notes: str = ""

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass
class IngestionResult:
    ledger: list[LedgerLine]
    profiles: list[DocumentProfile]


class DocumentIngestionError(ValueError):
    def __init__(self, message: str, profiles: list[DocumentProfile]):
        super().__init__(message)
        self.profiles = profiles


TEXT_EXTENSIONS = {".txt", ".md", ".json", ".xml", ".html", ".htm"}
TABLE_EXTENSIONS = {".xlsx", ".xls", ".csv", ".tsv"}


def _text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(float(value)) else None
    raw = _text(value)
    if not raw:
        return None
    negative = raw.startswith("(") and raw.endswith(")")
    cleaned = raw.strip("() ").replace(",", "")
    # Accept common currency labels and accounting whitespace, but do not
    # turn arbitrary text containing digits (for example a cover date or an
    # account name with a number) into a financial amount.
    cleaned = re.sub(r"(?i)\b(?:aed|dh|dirhams?|usd|eur|gbp)\b", "", cleaned)
    cleaned = cleaned.replace("−", "-").strip()
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", cleaned):
        return None
    try:
        result = float(cleaned)
    except ValueError:
        return None
    return -abs(result) if negative else result


def _read_docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    parts: list[str] = []
    for paragraph in root.iter():
        if paragraph.tag.endswith("}p"):
            words = [node.text or "" for node in paragraph.iter() if node.tag.endswith("}t")]
            if words:
                parts.append("".join(words).strip())
    return "\n".join(part for part in parts if part)


def _read_text(path: Path) -> str:
    if path.suffix.lower() == ".json":
        try:
            return json.dumps(json.loads(path.read_text(encoding="utf-8", errors="ignore")), indent=2)
        except json.JSONDecodeError:
            pass
    if path.suffix.lower() == ".docx":
        return _read_docx_text(path)
    if path.suffix.lower() == ".pdf":
        try:
            import pdfplumber

            with pdfplumber.open(path) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        except Exception as exc:  # OCR and malformed PDFs become review items.
            return f"[PDF extraction unavailable: {exc}]"
    return path.read_text(encoding="utf-8", errors="ignore")


def _candidate_header(rows: list[list[Any]]) -> tuple[int, list[str]]:
    limit = min(25, len(rows))
    semantic_tokens = (
        "account", "description", "particular", "item", "name", "ledger",
        "amount", "balance", "debit", "credit", "total", "value", "revenue",
    )
    best_index, best_score, best_headers = 0, float("-inf"), []
    for index in range(limit):
        headers = [_text(value).lower() for value in rows[index]]
        nonempty = [value for value in headers if value]
        if not nonempty:
            continue
        semantic = sum(1 for value in nonempty if any(token in value for token in semantic_tokens))
        text_count = sum(1 for value in nonempty if _number(value) is None)
        # Real client workbooks often place a title and an explanatory
        # sentence above the table. Prefer rows that have a table-shaped
        # block below them instead of treating a one-cell title containing
        # words such as "amounts" as the header.
        populated_width = len(nonempty)
        following_rows = rows[index + 1 : index + 6]
        table_shaped_rows = sum(
            sum(bool(_text(value)) for value in row) >= 2
            for row in following_rows
        )
        score = (
            semantic * 5
            + text_count
            + populated_width * 2
            + table_shaped_rows * 2
            - index * 0.05
        )
        if score > best_score:
            best_index, best_score, best_headers = index, score, headers
    return best_index, best_headers


def _discover_rows(rows: list[list[Any]]) -> tuple[list[tuple[str, float, int]], str]:
    indexed_rows = [
        (row_number, row)
        for row_number, row in enumerate(rows, start=1)
        if any(_text(value) for value in row)
    ]
    if len(indexed_rows) < 2:
        return [], "No tabular rows detected."

    rows = [row for _, row in indexed_rows]
    width = max(len(row) for row in rows)

    name_tokens = ("account", "description", "particular", "item", "name", "ledger", "narration")
    amount_tokens = ("amount", "balance", "debit", "credit", "value", "total", "net", "closing")
    name_index = None
    name_scores: list[tuple[float, int]] = []
    amount_scores: list[tuple[float, int]] = []
    amount_header_signals: dict[int, bool] = {}
    for column in range(width):
        values = [row[column] if column < len(row) else None for row in rows]
        nonempty = [_text(value) for value in values if _text(value)]
        numeric = [_number(value) for value in values if _text(value)]
        if not nonempty:
            continue
        numeric_ratio = sum(value is not None for value in numeric) / len(nonempty)

        # Header labels may occur several rows into a workbook and there may
        # be more than one table per sheet. Collect signals from table-shaped
        # rows in the first part of the sheet instead of relying on one global
        # header row.
        name_hint = 0
        amount_hint = 0
        period_hint = 0
        for row in rows[: min(40, len(rows))]:
            populated = [_text(cell) for cell in row if _text(cell)]
            cell = _text(row[column] if column < len(row) else "").lower()
            # Do not treat ordinary data rows such as "Total ... 4,800,000"
            # as headers merely because their description contains a header
            # word. Header rows normally contain labels/periods, not parsed
            # numeric amounts.
            if (
                len(populated) < 2
                or not cell
                or any(_number(value) is not None for value in row)
            ):
                continue
            if any(token in cell for token in name_tokens):
                name_hint += 1
            if any(token in cell for token in amount_tokens):
                amount_hint += 1
            if re.search(r"(?:19|20)\d{2}", cell):
                period_hint += 1
            if cell in {"aed", "dr", "cr", "debit", "credit"}:
                amount_hint += 1

        name_scores.append((name_hint * 6 + (1 - numeric_ratio) * 4 - column * 0.001, column))
        amount_header_signals[column] = bool(amount_hint or period_hint)
        amount_scores.append(
            (
                amount_hint * 7 + period_hint * 6 + numeric_ratio * 5 - column * 0.001,
                column,
            )
        )

    if name_scores:
        # Prefer a mostly textual column for the account/description. This
        # avoids mistaking trial-balance account numbers or note numbers for
        # the financial amount column.
        name_index = max(name_scores)[1]

    # Only columns with at least one parseable number can provide amounts.
    # Never call max() on an empty filtered generator: cover, notes, and
    # narrative sheets are valid workbook content but are not ledgers.
    amount_candidates = [
        (score, column)
        for score, column in amount_scores
        if column != name_index
        and (amount_header_signals[column] or column > name_index)
        and any(
            _number(row[column] if column < len(row) else None) is not None for row in rows
        )
    ]
    if name_index is None or not amount_candidates:
        return [], "Could not identify a text column and an amount column."

    # Prefer the first amount column on a tie (usually the current period),
    # while allowing header hints to win when a workbook labels debit/credit
    # or amount columns explicitly.
    amount_candidates.sort(key=lambda item: (item[0], -item[1]), reverse=True)
    amount_indices = [column for _, column in amount_candidates]

    parsed: list[tuple[str, float, int]] = []
    for row_index, row in indexed_rows:
        account = _text(row[name_index] if name_index < len(row) else "")
        amount = next(
            (
                _number(row[column] if column < len(row) else None)
                for column in amount_indices
                if _number(row[column] if column < len(row) else None) is not None
            ),
            None,
        )
        if account and amount is not None and account.lower() not in {"total", "subtotal", "grand total"}:
            parsed.append((account, amount, row_index))
    return parsed, f"Detected text column {name_index + 1} and amount column(s) {', '.join(str(column + 1) for column in amount_indices)}."


_METADATA_SHEET_NAMES = {
    "cover",
    "index",
    "contents",
    "table of contents",
    "toc",
    "readme",
    "instructions",
    "metadata",
}


def _is_metadata_sheet(sheet_name: str) -> bool:
    normalized = re.sub(r"\s+", " ", sheet_name.strip().lower())
    return normalized in _METADATA_SHEET_NAMES


def _spreadsheet_rows(path: Path) -> tuple[list[tuple[str, float, int, str]], list[str], str]:
    rows_found: list[tuple[str, float, int, str]] = []
    sheet_names: list[str] = []
    notes: list[str] = []
    if path.suffix.lower() in {".csv", ".tsv"}:
        separator = "\t" if path.suffix.lower() == ".tsv" else ","
        frame = pd.read_csv(path, sep=separator, header=None, dtype=object)
        parsed, note = _discover_rows(frame.values.tolist())
        return [(account, amount, row, "") for account, amount, row in parsed], [path.stem], note

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            sheet_names.append(sheet.title)
            if _is_metadata_sheet(sheet.title):
                notes.append(f"{sheet.title}: skipped metadata sheet")
                continue
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            try:
                parsed, note = _discover_rows(rows)
            except (TypeError, ValueError, ZeroDivisionError) as exc:
                # One malformed or non-tabular sheet must not prevent valid
                # tables in the rest of the workbook from being extracted.
                parsed = []
                note = f"skipped sheet after extraction error: {exc}"
            rows_found.extend((account, amount, row, sheet.title) for account, amount, row in parsed)
            notes.append(f"{sheet.title}: {note}")
    finally:
        workbook.close()
    return rows_found, sheet_names, "; ".join(notes) or "No financial table detected in workbook sheets."


def _text_rows(text: str) -> list[tuple[str, float, int]]:
    parsed: list[tuple[str, float, int]] = []
    amount_pattern = re.compile(r"^(.*?)(?:\s+|:)\(?([\d,]+(?:\.\d+)?)\)?\s*$")
    for line_number, line in enumerate(text.splitlines(), start=1):
        match = amount_pattern.match(line.strip())
        if not match:
            continue
        amount = _number(match.group(2))
        account = match.group(1).strip(" -:\t")
        if account and amount is not None:
            parsed.append((account, amount, line_number))
    return parsed


def inspect_document(path: Path) -> tuple[DocumentProfile, list[tuple[str, float, int, str]]]:
    extension = path.suffix.lower()
    if extension in TABLE_EXTENSIONS:
        try:
            rows, sheets, notes = _spreadsheet_rows(path)
            profile = DocumentProfile(path.name, extension, "structured_table", "extracted" if rows else "needs_review", len(rows), sheets, "", notes)
            return profile, rows
        except Exception as exc:
            profile = DocumentProfile(path.name, extension, "structured_table", "needs_review", notes=f"Table extraction failed: {exc}")
            return profile, []

    if extension in {".pdf", ".docx"} or extension in TEXT_EXTENSIONS:
        text = _read_text(path)
        rows = _text_rows(text)
        profile = DocumentProfile(
            path.name,
            extension,
            "document_text",
            "extracted" if rows else "needs_review",
            len(rows),
            [],
            text[:500],
            "Amount-bearing text detected." if rows else "Text extracted; no financial rows detected by the local fallback.",
        )
        return profile, [(account, amount, row, "") for account, amount, row in rows]

    return DocumentProfile(path.name, extension or "unknown", "binary_or_unknown", "needs_review", notes="Format accepted but requires a document-intelligence/OCR adapter."), []


def ingest_documents(paths: list[Path]) -> IngestionResult:
    profiles: list[DocumentProfile] = []
    ledger: list[LedgerLine] = []
    for path in paths:
        profile, rows = inspect_document(path)
        profiles.append(profile)
        for account, amount, source_row, source_sheet in rows:
            ledger.append(
                LedgerLine(
                    account_name=account,
                    amount=amount,
                    category=categorize(account),
                    description=account,
                    source_file=path.name,
                    source_sheet=source_sheet,
                    source_row=source_row,
                )
            )

    ledger = _deduplicate_repeated_source_rows(ledger)

    if not ledger:
        raise DocumentIngestionError(
            "No financial rows could be extracted from the uploaded documents. "
            "An AI document-understanding adapter or a human mapping is required for this file set.",
            profiles,
        )
    return IngestionResult(ledger=ledger, profiles=profiles)


def _deduplicate_repeated_source_rows(ledger: list[LedgerLine]) -> list[LedgerLine]:
    """Keep one canonical row when a workbook repeats the same fact.

    Financial statement workbooks commonly repeat dividend income in the P&L,
    notes, trial balance, and working papers. Those are distinct source rows
    but one economic item; counting all of them would multiply the tax
    adjustment. Dividend income is therefore deduplicated by document,
    category, and absolute amount, with the P&L preferred as the canonical
    provenance. Other categories use the normalized account label as an
    additional guard so unrelated transactions are not merged.
    """
    chosen: dict[tuple[str, str, float], tuple[int, LedgerLine]] = {}
    order: list[tuple[str, str, float]] = []
    sheet_priority = {"p&l": 0, "tb": 1, "working papers": 2, "notes": 3, "bs": 4, "oci": 5, "soce": 6}

    for position, line in enumerate(ledger):
        account = re.sub(r"[^a-z0-9 ]+", " ", line.account_name.lower())
        account = re.sub(r"\b(received|during the year|exempt income|uae|listed|portfolio)\b", " ", account)
        account = re.sub(r"\s+", " ", account).strip()
        if line.category.name == "DIVIDEND_INCOME":
            identity = "dividend_income"
        else:
            identity = account
        key = (line.source_file, identity, round(abs(line.amount), 2))
        priority = sheet_priority.get(line.source_sheet.lower(), 99)
        current = chosen.get(key)
        if current is None:
            chosen[key] = (priority, line)
            order.append(key)
        elif priority < current[0]:
            chosen[key] = (priority, line)

    return [chosen[key][1] for key in order]

"""Regression tests for irregular, multi-sheet client workbooks."""
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from app.intake.dynamic_documents import DocumentIngestionError, ingest_documents, inspect_document


class TestDynamicWorkbookIntake(unittest.TestCase):
    def _make_workbook(self, path: Path, include_ledger: bool = True) -> None:
        workbook = Workbook()
        cover = workbook.active
        cover.title = "Cover"
        cover["B8"] = "GULF HORIZON GENERAL TRADING L.L.C."
        cover["B9"] = "FINANCIAL STATEMENTS"

        if include_ledger:
            trial_balance = workbook.create_sheet("TB")
            trial_balance.append(["GULF HORIZON GENERAL TRADING L.L.C."])
            trial_balance.append(["TRIAL BALANCE"])
            trial_balance.append(["A/c", "Ledger account", "31.12.2025", "31.12.2024"])
            trial_balance.append([1, "Sales Revenue", 100000, 90000])
            trial_balance.append([2, "Client entertainment", 5000, 4000])

        workbook.save(path)

    def test_metadata_sheet_does_not_abort_valid_sheet_extraction(self):
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "irregular.xlsx"
            self._make_workbook(path)

            profile, rows = inspect_document(path)

        self.assertEqual(profile.extraction_status, "extracted")
        self.assertEqual(len(rows), 2)
        self.assertEqual({row[0] for row in rows}, {"Sales Revenue", "Client entertainment"})
        self.assertIn("Cover: skipped metadata sheet", profile.notes)

    def test_workbook_with_only_cover_returns_actionable_ingestion_error(self):
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "cover_only.xlsx"
            self._make_workbook(path, include_ledger=False)

            with self.assertRaises(DocumentIngestionError) as raised:
                ingest_documents([path])

        self.assertIn("No financial rows could be extracted", str(raised.exception))
        self.assertEqual(raised.exception.profiles[0].extraction_status, "needs_review")
        self.assertIn("skipped metadata sheet", raised.exception.profiles[0].notes)

    def test_repeated_dividend_fact_is_counted_once_and_keeps_canonical_sheet(self):
        with TemporaryDirectory() as temporary:
            path = Path(temporary) / "repeated_fact.xlsx"
            workbook = Workbook()
            workbook.remove(workbook.active)
            for title, label, amount in (
                ("P&L", "Dividend income - U.A.E. listed equity portfolio", 180000),
                ("Notes", "Dividend income - U.A.E. listed equity portfolio", 180000),
                ("TB", "Dividend income - U.A.E. listed equity portfolio", -180000),
                ("Working Papers", "Dividend income received (exempt income)", 180000),
            ):
                sheet = workbook.create_sheet(title)
                sheet.append(["Description", "Amount"])
                sheet.append([label, amount])
            workbook.save(path)

            result = ingest_documents([path])

        dividends = [line for line in result.ledger if "dividend income" in line.account_name.lower()]
        self.assertEqual(len(dividends), 1)
        self.assertEqual(dividends[0].source_sheet, "P&L")


if __name__ == "__main__":
    unittest.main()
